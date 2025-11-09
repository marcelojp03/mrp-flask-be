# controllers/report_ai_controller.py
from flask import Blueprint, request, jsonify, Response
from sqlalchemy import text
from app.db import db
from app.responses import Responses
import csv
import io
import re
import os
from datetime import datetime
from flask_jwt_extended import jwt_required, get_jwt

# Importaciones para exportación
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError:
    Workbook = None

try:
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.lib import colors
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
except ImportError:
    SimpleDocTemplate = None

# ==== Config LLM (usa variables de entorno) ====
try:
    from openai import OpenAI
    OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")
    LLM_MODEL = os.getenv("LLM_MODEL", "gpt-4o-mini")
    openai_client = OpenAI(api_key=OPENAI_API_KEY) if OPENAI_API_KEY else None
except ImportError:
    openai_client = None
    OPENAI_API_KEY = None

report_ai_bp = Blueprint("report_ai", __name__, url_prefix="/api/reports")

# ==== Schema dinámico: obtiene todas las tablas y columnas de la BD ====
def _get_database_schema(org_id: int):
    """
    Obtiene el esquema completo de la base de datos mediante introspección de SQLAlchemy
    Retorna dict con tablas y sus columnas CON TIPOS DE DATOS
    """
    from sqlalchemy import inspect
    inspector = inspect(db.engine)
    schema = {}
    tables_with_org = set()
    
    for table_name in inspector.get_table_names():
        columns = []
        has_org_id = False
        
        for column in inspector.get_columns(table_name):
            col_name = column['name']
            col_type = str(column['type'])
            
            # Incluir nombre y tipo de dato
            columns.append(f"{col_name}:{col_type}")
            
            if col_name == 'org_id':
                has_org_id = True
        
        schema[table_name] = columns
        if has_org_id:
            tables_with_org.add(table_name)
    
    return schema, tables_with_org

# Cache del esquema (se actualiza cada vez que se llama al endpoint)
WHITELIST = {}
TABLES_WITH_ORG = set()

FORBIDDEN = re.compile(r"\b(DROP|DELETE|TRUNCATE|UPDATE|INSERT|ALTER|EXEC|EXECUTE)\b", re.I)

def _schema_str():
    lines = []
    for t, cols in WHITELIST.items():
        lines.append(f"- {t}({', '.join(cols)})")
    return "\n".join(lines)

def _normalize_sql(s: str) -> str:
    # Quita fences y tags
    s = re.sub(r"```(?:sql)?\s*|\s*```", "", s, flags=re.I).strip()
    # Sin punto y coma final
    s = s.rstrip("; \n\t")
    return s

def _has_only_select_single_stmt(sql: str) -> bool:
    if FORBIDDEN.search(sql): 
        return False
    # no múltiple sentencia
    if ";" in sql: 
        return False
    # solo SELECT
    return sql.strip().upper().startswith("SELECT")

def _contains_org_predicate(sql: str) -> bool:
    return re.search(r"\borg_id\b", sql, re.I) is not None

def _inject_limit(sql: str, limit: int) -> str:
    if re.search(r"\bLIMIT\s+\d+", sql, re.I):
        return sql
    return f"{sql}\nLIMIT {limit}"

def _llm_generate_sql(nl_query: str, org_id: int, default_limit: int):
    if not openai_client:
        raise RuntimeError("OpenAI no configurado (falta OPENAI_API_KEY)")
    
    # Actualizar esquema dinámicamente
    global WHITELIST, TABLES_WITH_ORG
    WHITELIST, TABLES_WITH_ORG = _get_database_schema(org_id)
    
    system = (
        "Eres un generador de SQL para PostgreSQL. Devuelves SOLO una sentencia SELECT válida y nada más. "
        "Debes respetar EXACTAMENTE los tipos de datos de cada columna. "
        "IMPORTANTE: Prioriza columnas con información legible (nombres, códigos, descripciones) sobre IDs numéricos."
    )
    user = f"""
Base de datos PostgreSQL (formato: columna:tipo):
{_schema_str()}

Reglas OBLIGATORIAS:
1) SOLO una sentencia SELECT, sin comentarios, sin backticks, sin 'sql'.
2) NUNCA uses DROP/DELETE/UPDATE/INSERT/ALTER/EXEC.
3) Incluye siempre LIMIT {default_limit} si el usuario no especifica otro límite.
4) Cuando consultes tablas con org_id ({', '.join(sorted(TABLES_WITH_ORG))}), debes filtrar por org_id = {org_id}.
5) Usa nombres EXACTOS de tablas/columnas (sensible a errores).
6) Respeta los TIPOS DE DATOS:
   - BOOLEAN: usa TRUE/FALSE (no 'active', 'true', 1, 0)
   - INTEGER/BIGINT: números sin comillas
   - VARCHAR/TEXT: texto entre comillas simples
   - TIMESTAMP: formato '2025-01-01 12:00:00'
7) PRIORIZA INFORMACIÓN LEGIBLE:
   - Selecciona columnas 'name', 'code', 'description' en lugar de solo 'id'
   - En JOIN, incluye nombres de tablas relacionadas, no solo foreign keys
   - Evita SELECT * a menos que sea explícitamente solicitado
   - El usuario NO necesita ver IDs a menos que los solicite específicamente
8) PostgreSQL puro.

Pregunta del usuario:
\"\"\"{nl_query}\"\"\"

Devuelve únicamente el SQL (una sola línea o varias, pero una única sentencia).
"""
    r = openai_client.chat.completions.create(
        model=LLM_MODEL,
        messages=[{"role":"system","content":system},{"role":"user","content":user}],
        temperature=0.1,
    )
    sql = _normalize_sql(r.choices[0].message.content or "")
    return sql

def _llm_fix_sql(bad_sql: str, err: str, org_id: int, default_limit: int):
    if not openai_client:
        raise RuntimeError("OpenAI no configurado")
    
    system = (
        "Eres un corrector de SQL PostgreSQL. Devuelves SOLO una sentencia SELECT válida. "
        "Debes respetar EXACTAMENTE los tipos de datos de cada columna. "
        "Prioriza columnas con información legible (nombres, códigos) sobre IDs."
    )
    user = f"""
Corrige la siguiente consulta para PostgreSQL (formato columna:tipo):
{_schema_str()}

Reglas:
- SOLO una sentencia SELECT (sin ; final, sin backticks).
- Debe incluir LIMIT {default_limit} si no existe.
- Si consulta tablas con org_id ({', '.join(sorted(TABLES_WITH_ORG))}), filtra con org_id = {org_id}.
- Respeta los TIPOS DE DATOS:
  * BOOLEAN: usa TRUE/FALSE (no 'active', 'true', 1, 0)
  * INTEGER/BIGINT: números sin comillas
  * VARCHAR/TEXT: texto entre comillas simples
  * TIMESTAMP: formato '2025-01-01 12:00:00'
- INFORMACIÓN LEGIBLE:
  * Selecciona columnas 'name', 'code', 'description' en lugar de solo 'id'
  * En JOIN, incluye nombres de tablas relacionadas
  * El usuario NO necesita ver IDs a menos que los solicite

Consulta con error:
{bad_sql}

Error de PostgreSQL:
{err}

Devuelve únicamente el SQL corregido.
"""
    r = openai_client.chat.completions.create(
        model=LLM_MODEL,
        messages=[{"role":"system","content":system},{"role":"user","content":user}],
        temperature=0.1,
    )
    sql = _normalize_sql(r.choices[0].message.content or "")
    return sql

def _execute_readonly(sql: str, params: dict, as_csv=False):
    # statement_timeout opcional
    db.session.execute(text("SET LOCAL statement_timeout = '8s'"))
    result = db.session.execute(text(sql), params)
    cols = [c for c in result.keys()]
    rows = result.fetchall()

    if as_csv:
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(cols)
        for r in rows:
            writer.writerow([_jsonify_cell(x) for x in r])
        output.seek(0)
        return output.getvalue(), cols, rows

    # JSON
    data = [dict(zip(cols, [_jsonify_cell(x) for x in r])) for r in rows]
    return data, cols, rows

def _jsonify_cell(x):
    if isinstance(x, datetime):
        return x.isoformat()
    if isinstance(x, bool):
        return x
    # Decimal→float
    if hasattr(x, "as_tuple"):
        return float(x)
    return x

def _audit(user_id: int, org_id: int, nl: str, sql: str, ok: bool, err: str, rowcount: int, ms: int):
    try:
        from app.models.system_log import SystemLog
        log = SystemLog(
            user_id=user_id,
            org_id=org_id,
            action=f"REPORT_NL ok={ok} rows={rowcount} err={err[:120] if err else ''}",
            path="/api/reports/nl",
            method="POST",
            ip=request.remote_addr or ""
        )
        db.session.add(log)
        db.session.commit()
    except:
        db.session.rollback()

# ==== Interpretación de resultados ====
def _llm_interpret_results(nl_query: str, sql: str, columns: list, rows: list, max_rows: int = 5):
    """Genera una interpretación en lenguaje natural de los resultados"""
    if not openai_client:
        return None
    
    # Limitar filas mostradas a la IA
    sample_rows = rows[:max_rows]
    total_rows = len(rows)
    
    system = "Eres un asistente que interpreta resultados de consultas SQL en lenguaje natural claro y conciso."
    user = f"""
Pregunta del usuario:
"{nl_query}"

SQL ejecutado:
{sql}

Resultados ({total_rows} fila(s) total, mostrando primeras {len(sample_rows)}):
Columnas: {', '.join(columns)}
Datos:
{chr(10).join([str(dict(zip(columns, [str(v) for v in row.values()] if isinstance(row, dict) else row))) for row in sample_rows])}

Genera una respuesta en lenguaje natural que:
1. Responda directamente la pregunta del usuario
2. Sea clara y concisa (máximo 3 oraciones)
3. Incluya números/datos relevantes
4. Use un tono profesional pero amigable

Respuesta:
"""
    
    try:
        r = openai_client.chat.completions.create(
            model=LLM_MODEL,
            messages=[{"role":"system","content":system},{"role":"user","content":user}],
            temperature=0.3,
            max_tokens=200
        )
        return r.choices[0].message.content.strip()
    except:
        return None

# ==== Exportación a Excel ====
def _generate_excel(columns: list, rows: list, title: str = "Reporte"):
    """Genera un archivo Excel con los datos"""
    if not Workbook:
        raise ImportError("openpyxl no está instalado")
    
    output = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"
    
    # Estilo para encabezados
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    # Escribir encabezados
    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    
    # Escribir datos
    for row_idx, row_data in enumerate(rows, 2):
        if isinstance(row_data, dict):
            for col_idx, col_name in enumerate(columns, 1):
                value = row_data.get(col_name, '')
                ws.cell(row=row_idx, column=col_idx, value=value)
        else:
            for col_idx, value in enumerate(row_data, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)
    
    # Ajustar ancho de columnas
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    wb.save(output)
    output.seek(0)
    return output.getvalue()

# ==== Exportación a PDF ====
def _generate_pdf(columns: list, rows: list, title: str = "Reporte", nl_query: str = ""):
    """Genera un archivo PDF con los datos"""
    if not SimpleDocTemplate:
        raise ImportError("reportlab no está instalado")
    
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=letter)
    elements = []
    styles = getSampleStyleSheet()
    
    # Título
    if nl_query:
        title_text = f"Reporte: {nl_query}"
    else:
        title_text = title
    
    title_para = Paragraph(f"<b>{title_text}</b>", styles['Title'])
    elements.append(title_para)
    elements.append(Spacer(1, 0.3*inch))
    
    # Fecha
    date_para = Paragraph(f"Generado: {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S')} UTC", styles['Normal'])
    elements.append(date_para)
    elements.append(Spacer(1, 0.2*inch))
    
    # Preparar datos para la tabla
    table_data = [columns]  # Encabezados
    
    for row_data in rows:
        if isinstance(row_data, dict):
            row = [str(row_data.get(col, '')) for col in columns]
        else:
            row = [str(v) for v in row_data]
        table_data.append(row)
    
    # Crear tabla
    table = Table(table_data)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#366092')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
    ]))
    
    elements.append(table)
    
    # Construir PDF
    doc.build(elements)
    output.seek(0)
    return output.getvalue()


@report_ai_bp.route("/nl", methods=["POST"])
@jwt_required()
def natural_report():
    """
    Reportes por lenguaje natural (IA)
    Body: {query, format=json|csv|excel|pdf, limit=100, dry_run=false}
    """
    t0 = datetime.utcnow()
    payload = request.get_json() or {}
    nl_query = payload.get("query", "").strip()
    fmt = (payload.get("format") or "json").lower()
    limit = int(payload.get("limit") or 100)
    dry = bool(payload.get("dry_run") or False)

    if not nl_query:
        return Responses.error("Falta 'query'", 400)
    
    # Validar formato
    valid_formats = ['json', 'csv', 'excel', 'xlsx', 'pdf']
    if fmt not in valid_formats:
        return Responses.error(f"Formato inválido. Use: {', '.join(valid_formats)}", 400)

    jwt_data = get_jwt()
    user_id = jwt_data.get("sub")
    org_id = jwt_data.get("org_id")
    
    if not org_id:
        return Responses.error("Token sin org_id", 401)

    # Rate-limit: validar límite diario de reportes IA según plan
    from app.services.saas_guard_service import SaasGuardService
    try:
        saas_guard = SaasGuardService()
        saas_guard.assert_can_ai_reports_today(org_id)
    except ValueError as e:
        return Responses.error(str(e), 429)

    if not openai_client:
        return Responses.error("OpenAI no configurado en el servidor", 500)

    try:
        # 1) Generar SQL
        sql = _llm_generate_sql(nl_query, org_id, limit)
        sql = _inject_limit(sql, limit)

        # Validaciones
        if not _has_only_select_single_stmt(sql):
            return Responses.error("La IA no generó un SELECT válido", 400)

        # org guard
        if any(t in sql for t in TABLES_WITH_ORG) and not _contains_org_predicate(sql):
            return Responses.error("Falta filtro org_id en la consulta", 400)

        if dry:
            return Responses.success({"sql": sql}, "SQL generado (dry-run)")

        # 2) Ejecutar (read-only)
        params = {"org_id": org_id}
        as_csv = (fmt == "csv")
        data, cols, rows = _execute_readonly(sql, params, as_csv=as_csv)

        ms = (datetime.utcnow() - t0).total_seconds() * 1000
        _audit(user_id, org_id, nl_query, sql, True, None, len(rows), int(ms))

        # 3) Retornar según formato
        if fmt == "csv":
            return Response(
                data,
                mimetype="text/csv",
                headers={"Content-Disposition": f"attachment; filename=report_{datetime.utcnow().date()}.csv"}
            )
        elif fmt in ["excel", "xlsx"]:
            try:
                excel_data = _generate_excel(cols, data, title=f"Reporte: {nl_query[:50]}")
                return Response(
                    excel_data,
                    mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename=report_{datetime.utcnow().date()}.xlsx"}
                )
            except ImportError:
                return Responses.error("Exportación a Excel no disponible (openpyxl no instalado)", 500)
        elif fmt == "pdf":
            try:
                pdf_data = _generate_pdf(cols, data, title="Reporte", nl_query=nl_query)
                return Response(
                    pdf_data,
                    mimetype="application/pdf",
                    headers={"Content-Disposition": f"attachment; filename=report_{datetime.utcnow().date()}.pdf"}
                )
            except ImportError:
                return Responses.error("Exportación a PDF no disponible (reportlab no instalado)", 500)
        else:
            # JSON (default)
            # Generar interpretación en lenguaje natural
            interpretation = _llm_interpret_results(nl_query, sql, cols, data)
            
            return Responses.success({
                "sql": sql,
                "columns": cols,
                "rows": data,
                "interpretation": interpretation,
                "summary": {
                    "total_rows": len(rows),
                    "execution_time_ms": int(ms)
                },
                "export_options": ["json", "csv", "excel", "pdf"]
            }, "OK")

    except Exception as ex:
        # intento de autocorrección
        err = str(ex)
        try:
            fix = _llm_fix_sql(sql if 'sql' in locals() else '', err, org_id, limit)
            if not _has_only_select_single_stmt(fix):
                raise ValueError("Corrección inválida")
            fix = _inject_limit(fix, limit)
            params = {"org_id": org_id}
            data, cols, rows = _execute_readonly(fix, params, as_csv=(fmt == "csv"))

            ms = (datetime.utcnow() - t0).total_seconds() * 1000
            _audit(user_id, org_id, nl_query, fix, True, None, len(rows), int(ms))

            # Retornar según formato (auto-fix)
            if fmt == "csv":
                return Response(
                    data, mimetype="text/csv",
                    headers={"Content-Disposition": f"attachment; filename=report_{datetime.utcnow().date()}.csv"}
                )
            elif fmt in ["excel", "xlsx"]:
                excel_data = _generate_excel(cols, data, title=f"Reporte: {nl_query[:50]}")
                return Response(
                    excel_data,
                    mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename=report_{datetime.utcnow().date()}.xlsx"}
                )
            elif fmt == "pdf":
                pdf_data = _generate_pdf(cols, data, title="Reporte", nl_query=nl_query)
                return Response(
                    pdf_data,
                    mimetype="application/pdf",
                    headers={"Content-Disposition": f"attachment; filename=report_{datetime.utcnow().date()}.pdf"}
                )
            else:
                # JSON
                # Generar interpretación para auto-fix también
                interpretation = _llm_interpret_results(nl_query, fix, cols, data)
                
                return Responses.success({
                    "sql": fix,
                    "columns": cols,
                    "rows": data,
                    "interpretation": interpretation,
                    "summary": {
                        "total_rows": len(rows),
                        "execution_time_ms": int(ms)
                    },
                    "export_options": ["json", "csv", "excel", "pdf"]
                }, "OK (auto-fix)")

        except Exception as ex2:
            ms = (datetime.utcnow() - t0).total_seconds() * 1000
            _audit(user_id, org_id, nl_query, (locals().get('sql') or ''), False, f"{err} | {ex2}", 0, int(ms))
            return Responses.error(f"No se pudo ejecutar el reporte: {err}", 400)
